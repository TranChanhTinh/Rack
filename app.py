from flask import Flask, render_template, request, jsonify, send_file
from pathlib import Path
import json, uuid, io, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)
DATA_DIR = Path(os.environ.get('DATA_DIR', str(Path(__file__).resolve().parent)))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DATA_FILE = DATA_DIR / 'khorack_data.json'
RACKS = [
    {'id': 'ML', 'name': 'KHU 1', 'color': '#1675e8', 'tiers': ['T3', 'T2', 'T1']},
    {'id': 'IV1', 'name': 'KHU 2', 'color': '#169d70', 'tiers': ['T3', 'T2', 'T1']},
    {'id': 'IV2', 'name': 'KHU 3', 'color': '#f28a20', 'tiers': ['T3', 'T2', 'T1']},
    {'id': 'IV3', 'name': 'KHU 4', 'color': '#6550cf', 'tiers': ['T1']},
]

def default_state():
    return {'tabs': [{'id': 'default', 'name': 'Kho chính', 'rackNames': {r['id']: r['name'] for r in RACKS}, 'boxes': {}, 'order': {}}], 'activeTab': 'default'}

def load_state():
    try:
        raw = json.loads(DATA_FILE.read_text(encoding='utf-8'))
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return default_state()
    # Migrate legacy flat box dictionary to the new tab format.
    if isinstance(raw, dict) and 'tabs' not in raw:
        state = default_state()
        state['tabs'][0]['boxes'] = {k: v for k, v in raw.items() if isinstance(k, str) and k.count('|') == 2}
        return state
    if not isinstance(raw, dict) or not isinstance(raw.get('tabs'), list) or not raw['tabs']:
        return default_state()
    old_defaults = {'ML': 'ML', 'IV1': 'IV#1', 'IV2': 'IV#2', 'IV3': 'IV#3'}
    new_defaults = {r['id']: r['name'] for r in RACKS}
    for tab in raw['tabs']:
        tab.setdefault('rackNames', dict(new_defaults))
        # Update only the old built-in names; keep names the user deliberately customized.
        for rack_id, old_name in old_defaults.items():
            if tab['rackNames'].get(rack_id) in (None, '', old_name):
                tab['rackNames'][rack_id] = new_defaults[rack_id]
        tab.setdefault('boxes', {})
        tab.setdefault('order', {})
    if not any(t.get('id') == raw.get('activeTab') for t in raw['tabs']):
        raw['activeTab'] = raw['tabs'][0]['id']
    return raw

def save_state(state):
    DATA_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding='utf-8')

@app.get('/')
def index():
    return render_template('index.html', racks=RACKS, state=load_state())

@app.get('/excel-template')
def excel_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'BoxData'
    headers = ['Tab', 'Rack', 'Tier', 'Position', 'BoxCode']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176FCE')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['Kho chính', 'KHU 1', 'T3', 1, 'ML-101'])
    ws.append(['Kho chính', 'KHU 2', 'T2', 4, 'BOX-204'])
    ws.append(['', '', '', '', ''])
    ws['A5'] = 'Hướng dẫn: mỗi dòng là một box. Xóa 2 dòng ví dụ trước khi nhập dữ liệu thật.'
    ws['A6'] = 'Rack nhận KHU 1, KHU 2, KHU 3, KHU 4 hoặc tên rack đã đổi (cũng chấp nhận mã cũ ML, IV#1, IV#2, IV#3). Position là số ô 1-24; BoxCode để trống nghĩa là ô trống.'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:E3'
    widths = [24, 18, 14, 14, 28]
    for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='KhoRack_import_template.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.get('/api/export-excel')
def export_excel():
    """Export the current rack layout from saved server state to a management workbook."""
    state = load_state()
    wb = Workbook()
    ws = wb.active
    ws.title = 'DanhSachBox'
    headers = ['Tab', 'Rack', 'Tier', 'Position', 'BoxCode', 'TrangThai']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176FCE')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for tab in state.get('tabs', []):
        boxes = tab.get('boxes', {})
        rack_names = tab.get('rackNames', {})
        for rack in RACKS:
            rack_name = rack_names.get(rack['id'], rack['name'])
            for tier in rack['tiers']:
                for pos in range(1, 25):
                    code = boxes.get(f"{rack['id']}|{tier}|{pos}", '')
                    ws.append([tab.get('name', 'Kho chính'), rack_name, tier, pos, code, 'Đã sử dụng' if code else 'Trống'])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col, width in enumerate([24, 20, 12, 12, 30, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=6):
        row[3].alignment = Alignment(horizontal='center')
        if row[4].value:
            row[4].font = Font(color='176FCE', bold=True)
    summary = wb.create_sheet('TongQuan')
    summary.append(['Tab', 'Rack', 'Số tầng', 'Tổng vị trí', 'Đã sử dụng', 'Còn trống'])
    for cell in summary[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176FCE')
    for tab in state.get('tabs', []):
        boxes = tab.get('boxes', {})
        for rack in RACKS:
            rack_name = tab.get('rackNames', {}).get(rack['id'], rack['name'])
            capacity = len(rack['tiers']) * 24
            used = sum(1 for tier in rack['tiers'] for pos in range(1, 25) if boxes.get(f"{rack['id']}|{tier}|{pos}", ''))
            summary.append([tab.get('name', 'Kho chính'), rack_name, len(rack['tiers']), capacity, used, capacity-used])
    for col, width in enumerate([24, 20, 12, 16, 16, 16], 1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.freeze_panes = 'A2'
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='KhoRack_DuLieu_QuanLy.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.post('/api/import-excel')
def import_excel():
    file = request.files.get('file')
    if not file or not file.filename.lower().endswith(('.xlsx', '.xlsm')):
        return jsonify({'ok': False, 'error': 'Vui lòng chọn file Excel .xlsx'}), 400
    try:
        wb = load_workbook(file, data_only=True)
        ws = wb['BoxData'] if 'BoxData' in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows: raise ValueError('File Excel không có dữ liệu')
        headers = [str(v or '').strip().lower() for v in rows[0]]
        aliases = {'tab': ['tab','kho','tên tab'], 'rack': ['rack','kệ','rack name'], 'tier': ['tier','tầng'], 'position': ['position','vị trí','ô'], 'boxcode': ['boxcode','box code','mã box','box','mã']}
        indexes = {}
        for key, names in aliases.items():
            for i, h in enumerate(headers):
                if h in names: indexes[key] = i; break
        if len(indexes) < 5: raise ValueError('Thiếu cột. Hãy dùng file mẫu Excel của KhoRack.')
        state = load_state(); imported = 0
        rack_ids = {r['id']: r['id'] for r in RACKS}
        for row in rows[1:]:
            if not row or all(v is None or str(v).strip()=='' for v in row): continue
            def val(k):
                i=indexes[k]; return row[i] if i < len(row) else None
            tab_name = str(val('tab') or 'Kho chính').strip()
            rack_text = str(val('rack') or '').strip()
            tier = str(val('tier') or '').strip().upper()
            pos_raw = val('position'); code = str(val('boxcode') or '').strip()
            if not rack_text or not tier or pos_raw is None: continue
            pos = int(pos_raw)
            if not 1 <= pos <= 24: raise ValueError(f'Position phải từ 1 đến 24 (nhận {pos})')
            rid = next((r['id'] for r in RACKS if rack_text.casefold() in (r['id'].casefold(), r['name'].casefold())), None)
            if rid is None:
                old_aliases = {'ML': 'ML', 'IV#1': 'IV1', 'IV#2': 'IV2', 'IV#3': 'IV3'}
                rid = next((v for k, v in old_aliases.items() if rack_text.casefold() == k.casefold()), None)
            if rid is None:
                rid = next((r['id'] for t in state['tabs'] for r in RACKS if t.get('rackNames',{}).get(r['id'],'').casefold()==rack_text.casefold()), None)
            if rid is None: raise ValueError(f'Rack không hợp lệ: {rack_text}')
            if tier not in next(r['tiers'] for r in RACKS if r['id']==rid): raise ValueError(f'Tầng {tier} không tồn tại trong rack {rack_text}')
            tab = next((t for t in state['tabs'] if t['name'].casefold()==tab_name.casefold()), None)
            if tab is None:
                tab={'id':'tab_'+uuid.uuid4().hex[:10], 'name':tab_name[:40], 'rackNames':{r['id']:r['name'] for r in RACKS}, 'boxes':{}, 'order':{}}
                state['tabs'].append(tab)
            key=f'{rid}|{tier}|{pos}'
            if code: tab['boxes'][key]=code[:100]
            else: tab['boxes'].pop(key, None)
            imported += 1
        save_state(state)
        return jsonify({'ok':True,'imported':imported,'tabs':len(state['tabs'])})
    except Exception as exc:
        return jsonify({'ok':False,'error':str(exc)}),400

@app.get('/api/state')
def get_state():
    return jsonify(load_state())

@app.post('/api/state')
def update_state():
    state = request.get_json(silent=True)
    if not isinstance(state, dict) or not isinstance(state.get('tabs'), list) or not state['tabs']:
        return jsonify({'ok': False, 'error': 'Dữ liệu tab không hợp lệ'}), 400
    try:
        save_state(state)
    except OSError as exc:
        return jsonify({'ok': False, 'error': str(exc)}), 500
    return jsonify({'ok': True})

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', '5000')), debug=False)
